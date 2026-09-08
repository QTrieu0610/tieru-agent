"""Sandboxed local filesystem, document, and foreground command tools."""

from __future__ import annotations

import csv
import io
import json
import locale
import os
import re
import shlex
import subprocess
import zipfile
from pathlib import Path
from typing import Any

from tieru.tools.registry import Tool

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_UNCOMPRESSED = 100 * 1024 * 1024
MAX_OUTPUT_CHARS = 100_000
MAX_SEARCH_FILES = 1000
MAX_SEARCH_BYTES = 4 * 1024 * 1024
MAX_SHELL_OUTPUT_BYTES = 20_000
TEXT_TYPES = {".txt", ".md", ".json", ".csv"}
OFFICE_TYPES = {".docx", ".pptx", ".xlsx"}
SKIP_SEARCH_DIRS = {".git", ".tieru", "__pycache__", ".pytest_cache", ".ruff_cache"}
BLOCKED_SHELL = {
    "del", "erase", "format", "mkfs", "move", "mv", "rd", "reboot", "remove-item",
    "ren", "rename-item", "rm", "rmdir", "shutdown", "start", "start-process",
}
SHELL_META = re.compile(r"(?:[;&|<>\r\n\x60]|\$\()")
SECRET_ENV = re.compile(
    r"(?:api_?key|authorization|cookie|credential|password|secret|token)", re.IGNORECASE
)


class WorkspacePolicyError(ValueError):
    pass


class UnsupportedDocumentError(ValueError):
    pass


def _json(**fields: Any) -> str:
    return json.dumps(fields, ensure_ascii=False, sort_keys=True)


def _decode(data: bytes) -> tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    encodings = ["utf-8"]
    preferred = locale.getpreferredencoding(False)
    if preferred.lower() not in {"utf-8", "utf8"}:
        encodings.append(preferred)
    encodings.append("cp1252")
    for encoding in encodings:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeError("file is binary or uses an unsupported text encoding")


def _bounded(text: str, max_chars: int) -> tuple[str, bool]:
    limit = max(1, min(int(max_chars), MAX_OUTPUT_CHARS))
    return text[:limit], len(text) > limit


class LocalComputer:
    """All local actions are rooted under one immutable resolved workspace."""

    def __init__(self, root: Path):
        self.root = root.resolve(strict=True)
        if not self.root.is_dir():
            raise WorkspacePolicyError("workspace root must be an existing directory")

    def _path(self, value: str = ".", *, must_exist: bool = True) -> Path:
        raw = Path(value or ".").expanduser()
        candidate = (
            raw.resolve(strict=False)
            if raw.is_absolute()
            else (self.root / raw).resolve(strict=False)
        )
        if candidate != self.root and self.root not in candidate.parents:
            raise WorkspacePolicyError("path is outside the workspace root")
        if must_exist and not candidate.exists():
            raise FileNotFoundError(f"workspace path does not exist: {value}")
        return candidate

    def _relative(self, path: Path) -> str:
        value = path.relative_to(self.root).as_posix()
        return value or "."

    def list(self, path: str = ".", max_entries: int = 200) -> str:
        target = self._path(path)
        if not target.is_dir():
            raise NotADirectoryError(f"workspace path is not a directory: {path}")
        limit = max(1, min(int(max_entries), 1000))
        items = sorted(target.iterdir(), key=lambda item: (not item.is_dir(), item.name.casefold()))
        entries = []
        for item in items[:limit]:
            try:
                if item.is_symlink():
                    kind, size = "symlink", None
                else:
                    stat = item.stat()
                    kind = "directory" if item.is_dir() else "file" if item.is_file() else "other"
                    size = stat.st_size if item.is_file() else None
            except OSError:
                kind, size = "unavailable", None
            entries.append({"path": self._relative(item), "type": kind, "size": size})
        return _json(path=self._relative(target), entries=entries, truncated=len(items) > limit)

    def read(self, path: str, max_chars: int = 50_000) -> str:
        target = self._path(path)
        if not target.is_file():
            raise IsADirectoryError(f"workspace path is not a file: {path}")
        if target.stat().st_size > MAX_FILE_BYTES:
            raise WorkspacePolicyError(f"file exceeds the {MAX_FILE_BYTES}-byte read limit")
        with target.open("rb") as handle:
            data = handle.read(MAX_FILE_BYTES + 1)
        if b"\x00" in data:
            raise UnsupportedDocumentError("binary files require document_read when supported")
        text, encoding = _decode(data)
        content, truncated = _bounded(text, max_chars)
        return _json(
            path=self._relative(target), content=content, encoding=encoding, truncated=truncated
        )

    def search(self, query: str, path: str = ".", max_results: int = 50) -> str:
        if not query:
            raise ValueError("query must not be empty")
        target = self._path(path)
        limit = max(1, min(int(max_results), 200))
        lowered = query.casefold()
        results: list[dict[str, Any]] = []
        scanned_files = 0
        scanned_bytes = 0

        def candidates():
            if target.is_file():
                yield target
                return
            for current, dirs, files in os.walk(target, followlinks=False):
                dirs[:] = sorted(name for name in dirs if name not in SKIP_SEARCH_DIRS)
                for name in sorted(files):
                    yield Path(current) / name

        for item in candidates():
            if len(results) >= limit or scanned_files >= MAX_SEARCH_FILES:
                break
            try:
                resolved = self._path(str(item))
                size = resolved.stat().st_size
            except (OSError, WorkspacePolicyError):
                continue
            scanned_files += 1
            relative = self._relative(resolved)
            if lowered in relative.casefold():
                results.append({"path": relative, "line": None, "text": ""})
                if len(results) >= limit:
                    break
            if size > MAX_FILE_BYTES or scanned_bytes + size > MAX_SEARCH_BYTES:
                continue
            try:
                with resolved.open("rb") as handle:
                    data = handle.read(MAX_FILE_BYTES + 1)
                if b"\x00" in data:
                    continue
                text, _encoding = _decode(data)
            except (OSError, UnicodeError):
                continue
            scanned_bytes += len(data)
            for line_number, line in enumerate(text.splitlines(), 1):
                if lowered in line.casefold():
                    results.append({"path": relative, "line": line_number, "text": line[:500]})
                    if len(results) >= limit:
                        break
        truncated = (
            len(results) >= limit
            or scanned_files >= MAX_SEARCH_FILES
            or scanned_bytes >= MAX_SEARCH_BYTES
        )
        candidate_paths = list(dict.fromkeys(item["path"] for item in results))
        return _json(
            query=query,
            path=self._relative(target),
            results=results,
            candidate_paths=candidate_paths,
            next_action=(
                {"tool": "filesystem_read", "path": candidate_paths[0]}
                if len(candidate_paths) == 1
                else None
            ),
            scanned_files=scanned_files,
            truncated=truncated,
        )

    def write(
        self,
        path: str,
        content: str,
        overwrite: bool = False,
        encoding: str = "",
    ) -> str:
        target = self._path(path, must_exist=False)
        if target == self.root:
            raise IsADirectoryError("filesystem_write requires a file path")
        if not target.parent.is_dir():
            raise FileNotFoundError("parent directory does not exist; use filesystem_mkdir first")
        existed = target.exists()
        if existed and not target.is_file():
            raise IsADirectoryError(f"workspace path is not a file: {path}")
        if existed and not overwrite:
            raise FileExistsError("target exists; set overwrite=true after explicit approval")
        selected = encoding
        if not selected and existed:
            with target.open("rb") as handle:
                _old, selected = _decode(handle.read(MAX_FILE_BYTES + 1))
        selected = selected or "utf-8"
        with target.open("w", encoding=selected, newline="") as handle:
            handle.write(content)
        return _json(
            path=self._relative(target),
            bytes=target.stat().st_size,
            encoding=selected,
            created=not existed,
            overwritten=existed,
            status="ok",
        )

    def mkdir(self, path: str, parents: bool = False) -> str:
        target = self._path(path, must_exist=False)
        existed = target.exists()
        if existed and not target.is_dir():
            raise FileExistsError(f"workspace path exists and is not a directory: {path}")
        target.mkdir(parents=parents, exist_ok=True)
        return _json(path=self._relative(target), created=not existed, status="ok")

    def document_read(self, path: str, max_chars: int = 50_000) -> str:
        target = self._path(path)
        if not target.is_file():
            raise IsADirectoryError(f"workspace path is not a file: {path}")
        if target.stat().st_size > MAX_DOCUMENT_BYTES:
            raise WorkspacePolicyError(f"document exceeds the {MAX_DOCUMENT_BYTES}-byte limit")
        suffix = target.suffix.lower()
        if suffix in TEXT_TYPES:
            text = self._read_text_document(target, suffix)
        elif suffix == ".pdf":
            text = self._read_pdf(target)
        elif suffix in OFFICE_TYPES:
            self._validate_office_archive(target)
            text = {
                ".docx": self._read_docx,
                ".pptx": self._read_pptx,
                ".xlsx": self._read_xlsx,
            }[suffix](target)
        else:
            raise UnsupportedDocumentError(
                f"unsupported document type '{suffix or '(none)'}'; "
                "supported: txt, md, json, csv, pdf, docx, pptx, xlsx"
            )
        if not text.strip():
            raise UnsupportedDocumentError(
                "document contains no extractable text; image-only/scanned content is unsupported"
            )
        content, truncated = _bounded(text, max_chars)
        return _json(
            path=self._relative(target),
            type=suffix.lstrip("."),
            content=content,
            truncated=truncated,
        )

    @staticmethod
    def _read_text_document(path: Path, suffix: str) -> str:
        with path.open("rb") as handle:
            data = handle.read(MAX_DOCUMENT_BYTES + 1)
        text, _encoding = _decode(data)
        if suffix == ".json":
            return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        if suffix == ".csv":
            return "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text)))
        return text

    @staticmethod
    def _read_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise UnsupportedDocumentError(
                "PDF support requires the optional 'tieru-agent[documents]' dependency"
            ) from exc
        with path.open("rb") as handle:
            reader = PdfReader(handle)
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)

    @staticmethod
    def _read_docx(path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise UnsupportedDocumentError(
                "DOCX support requires the optional 'tieru-agent[documents]' dependency"
            ) from exc
        document = Document(path)
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(lines)

    @staticmethod
    def _read_pptx(path: Path) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise UnsupportedDocumentError(
                "PPTX support requires the optional 'tieru-agent[documents]' dependency"
            ) from exc
        presentation = Presentation(path)
        lines = []
        for index, slide in enumerate(presentation.slides, 1):
            values = [
                str(shape.text)
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and str(shape.text).strip()
            ]
            if values:
                lines.append(f"[Slide {index}]\n" + "\n".join(values))
        return "\n\n".join(lines)

    @staticmethod
    def _read_xlsx(path: Path) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise UnsupportedDocumentError(
                "XLSX support requires the optional 'tieru-agent[documents]' dependency"
            ) from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            lines = []
            for sheet in workbook.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    lines.append("\t".join("" if value is None else str(value) for value in row))
            return "\n".join(lines)
        finally:
            workbook.close()

    @staticmethod
    def _validate_office_archive(path: Path) -> None:
        try:
            with zipfile.ZipFile(path) as archive:
                total = sum(item.file_size for item in archive.infolist())
        except zipfile.BadZipFile as exc:
            raise UnsupportedDocumentError("invalid Office document archive") from exc
        if total > MAX_DOCUMENT_UNCOMPRESSED:
            raise WorkspacePolicyError("Office document exceeds the uncompressed content limit")

    def shell_run(self, command: str, timeout_seconds: int = 20) -> str:
        if not command.strip():
            raise ValueError("command must not be empty")
        if SHELL_META.search(command):
            raise WorkspacePolicyError(
                "shell chaining, redirection, substitution, and background syntax are blocked"
            )
        try:
            args = shlex.split(command, posix=os.name != "nt")
        except ValueError as exc:
            raise WorkspacePolicyError(f"invalid command quoting: {exc}") from exc
        if not args:
            raise ValueError("command must not be empty")
        args = [
            item[1:-1]
            if len(item) >= 2 and item[0] == item[-1] and item[0] in {"'", '"'}
            else item
            for item in args
        ]
        executable = Path(args[0].strip('"')).name.casefold()
        if executable in BLOCKED_SHELL:
            raise WorkspacePolicyError(f"destructive or background command is blocked: {executable}")
        lowered = [item.strip('"').casefold() for item in args]
        if self._is_install_command(lowered):
            raise WorkspacePolicyError("package installation is unavailable in shell_run")
        if executable == "git" and self._unsafe_git(lowered[1:]):
            raise WorkspacePolicyError(
                "Git mutation is unavailable in shell_run; use dedicated read-only Git tools"
            )
        timeout = max(1, min(int(timeout_seconds), 30))
        environment = {
            key: value for key, value in os.environ.items() if not SECRET_ENV.search(key)
        }
        try:
            completed = subprocess.run(
                args,
                cwd=self.root,
                env=environment,
                stdin=subprocess.DEVNULL,
                capture_output=True,
                timeout=timeout,
                check=False,
                shell=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise TimeoutError(f"shell command timed out after {timeout} seconds") from exc
        stdout_bytes = completed.stdout[:MAX_SHELL_OUTPUT_BYTES]
        remaining = max(0, MAX_SHELL_OUTPUT_BYTES - len(stdout_bytes))
        stderr_bytes = completed.stderr[:remaining]
        return _json(
            command=command,
            cwd=".",
            exit_code=completed.returncode,
            stdout=stdout_bytes.decode("utf-8", errors="replace"),
            stderr=stderr_bytes.decode("utf-8", errors="replace"),
            truncated=len(completed.stdout) + len(completed.stderr) > MAX_SHELL_OUTPUT_BYTES,
            status="ok" if completed.returncode == 0 else "error",
        )

    @staticmethod
    def _is_install_command(args: list[str]) -> bool:
        if not args:
            return False
        executable = args[0]
        managers = {"pip", "pip3", "uv", "npm", "pnpm", "yarn", "apt", "apt-get", "winget",
                    "choco", "brew"}
        if executable in managers:
            return any(item in {"add", "install"} for item in args[1:]) or executable in {
                "apt", "apt-get", "winget", "choco", "brew"
            }
        return (
            executable in {"python", "python3", "py"}
            and len(args) >= 4
            and args[1:4] in (["-m", "pip", "install"], ["-m", "uv", "add"])
        )

    @staticmethod
    def _unsafe_git(args: list[str]) -> bool:
        if not args or args == ["--version"]:
            return False
        return args[0] not in {"status", "diff", "log", "show", "rev-parse", "ls-files", "grep"}


def _schema(properties: dict | None = None, required: list[str] | None = None) -> dict:
    result = {
        "type": "object",
        "properties": properties or {},
        "additionalProperties": False,
    }
    if required:
        result["required"] = required
    return result


def make_tools(computer: LocalComputer) -> list[Tool]:
    path = {"type": "string", "maxLength": 4096}
    read_common = {
        "read_only": True,
        "default_policy": "allow",
        "risk": "low",
        "capabilities": ("filesystem.read",),
        "target_arg": "path",
        "scope": "workspace",
    }
    write_common = {
        "read_only": False,
        "default_policy": "confirm",
        "risk": "medium",
        "capabilities": ("filesystem.write",),
        "target_arg": "path",
        "scope": "workspace",
    }
    return [
        Tool(
            "filesystem_list",
            "List bounded entries under a directory in the current workspace only.",
            _schema(
                {
                    "path": path,
                    "max_entries": {"type": "integer", "minimum": 1, "maximum": 1000},
                }
            ),
            computer.list,
            operation="list",
            resource_type="filesystem",
            **read_common,
        ),
        Tool(
            "filesystem_read",
            "Read one bounded text file in the current workspace; binary files are rejected.",
            _schema(
                {
                    "path": path,
                    "max_chars": {"type": "integer", "minimum": 1, "maximum": MAX_OUTPUT_CHARS},
                },
                ["path"],
            ),
            computer.read,
            operation="read",
            resource_type="filesystem",
            **read_common,
        ),
        Tool(
            "filesystem_search",
            "Find candidate files by bounded filename/text search. Read the selected file before answering.",
            _schema(
                {
                    "query": {"type": "string", "minLength": 1, "maxLength": 500},
                    "path": path,
                    "max_results": {"type": "integer", "minimum": 1, "maximum": 200},
                },
                ["query"],
            ),
            computer.search,
            operation="search",
            resource_type="filesystem",
            **read_common,
        ),
        Tool(
            "filesystem_write",
            "Create or edit one text file inside the workspace. When editing an existing file, "
            "send overwrite=true; the call still remains subject to Trust. Never deletes or moves.",
            _schema(
                {
                    "path": path,
                    "content": {"type": "string", "maxLength": MAX_OUTPUT_CHARS},
                    "overwrite": {
                        "type": "boolean",
                        "description": (
                            "Set true when the target already exists and the requested task "
                            "explicitly authorizes editing it; false creates a new file only."
                        ),
                    },
                    "encoding": {
                        "type": "string",
                        "enum": ["", "utf-8", "utf-8-sig", "utf-16", "cp1252"],
                    },
                },
                ["path", "content"],
            ),
            computer.write,
            sensitive_args=("content",),
            operation="write",
            resource_type="filesystem",
            **write_common,
        ),
        Tool(
            "filesystem_mkdir",
            "Create one directory inside the workspace; never deletes or moves.",
            _schema({"path": path, "parents": {"type": "boolean"}}, ["path"]),
            computer.mkdir,
            operation="mkdir",
            resource_type="filesystem",
            **write_common,
        ),
        Tool(
            "document_read",
            "Extract bounded text from txt, md, json, csv, pdf, docx, pptx, or xlsx in the workspace.",
            _schema(
                {
                    "path": path,
                    "max_chars": {"type": "integer", "minimum": 1, "maximum": MAX_OUTPUT_CHARS},
                },
                ["path"],
            ),
            computer.document_read,
            operation="read_document",
            resource_type="document",
            **read_common,
        ),
        Tool(
            "shell_run",
            "Run one foreground command in the workspace. No chaining, background work, installs, or destructive commands.",
            _schema(
                {
                    "command": {"type": "string", "minLength": 1, "maxLength": 4000},
                    "timeout_seconds": {"type": "integer", "minimum": 1, "maximum": 30},
                },
                ["command"],
            ),
            computer.shell_run,
            read_only=False,
            default_policy="confirm",
            risk="high",
            capabilities=("process.execute",),
            operation="run",
            target_arg="command",
            fixed_target="workspace foreground process",
            scope="workspace",
            resource_type="process",
            reversible=False,
        ),
    ]
