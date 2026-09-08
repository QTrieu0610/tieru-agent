"""Deterministic M8 local-computer sandbox, parser, shell, and Replay contracts."""

from __future__ import annotations

import json
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest

from evals.helpers import ScriptedClient, make_waku, response, text_block, tool_block
from tieru.tools.computer import LocalComputer, make_tools
from tieru.tools.registry import ToolRegistry


def registry_for(root: Path, *, approve=True) -> tuple[LocalComputer, ToolRegistry]:
    computer = LocalComputer(root)
    registry = ToolRegistry(
        approval_handler=lambda _request: approve,
        trust_context={"base_path": str(root), "path_aliases": {"workspace": str(root)}},
    )
    for tool in make_tools(computer):
        registry.register(tool)
    return computer, registry


def parsed(registry: ToolRegistry, name: str, args: dict) -> dict:
    return json.loads(registry.execute(name, args))


def write_text_pdf(path: Path, text: str = "PDF fixture text") -> None:
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    )
    path.write_bytes(output)


def make_documents(root: Path) -> dict[str, str]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    expected = {
        "fixture.pdf": "PDF fixture text",
        "fixture.docx": "DOCX fixture text",
        "fixture.pptx": "PPTX fixture text",
        "fixture.xlsx": "XLSX fixture text",
    }
    write_text_pdf(root / "fixture.pdf")
    document = Document()
    document.add_paragraph(expected["fixture.docx"])
    document.save(root / "fixture.docx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(10, 10, 500, 100)
    box.text_frame.text = expected["fixture.pptx"]
    presentation.save(root / "fixture.pptx")
    workbook = Workbook()
    workbook.active["A1"] = expected["fixture.xlsx"]
    workbook.save(root / "fixture.xlsx")
    workbook.close()
    return expected


def test_filesystem_list_read_and_search(tmp_path):
    (tmp_path / "docs").mkdir()
    (tmp_path / "docs" / "answer.md").write_text(
        "Tieru local computer answer is forty-two.", encoding="utf-8"
    )
    _computer, registry = registry_for(tmp_path)

    listed = parsed(registry, "filesystem_list", {"path": "docs"})
    read = parsed(registry, "filesystem_read", {"path": "docs/answer.md"})
    searched = parsed(registry, "filesystem_search", {"query": "forty-two"})

    assert listed["entries"][0]["path"] == "docs/answer.md"
    assert read["encoding"] == "utf-8" and "forty-two" in read["content"]
    assert searched["results"][0]["path"] == "docs/answer.md"


@pytest.mark.parametrize("path", ["../outside.txt", "../../outside.txt"])
def test_path_traversal_is_blocked(path, tmp_path):
    (tmp_path.parent / "outside.txt").write_text("outside secret", encoding="utf-8")
    _computer, registry = registry_for(tmp_path)

    result = parsed(registry, "filesystem_read", {"path": path})

    assert result["error"]["code"] == "tool_execution_error"
    assert "outside secret" not in json.dumps(result)


def test_filesystem_write_requires_permission_and_preserves_encoding(tmp_path):
    existing = tmp_path / "legacy.txt"
    existing.write_text("café", encoding="cp1252")
    denied_computer, denied = registry_for(tmp_path, approve=False)

    result = parsed(denied, "filesystem_write", {"path": "new.txt", "content": "blocked"})
    denied_directory = parsed(denied, "filesystem_mkdir", {"path": "blocked-dir"})

    assert result["error"]["code"] == "tool_permission_denied"
    assert denied_directory["error"]["code"] == "tool_permission_denied"
    assert not (tmp_path / "new.txt").exists()
    assert not (tmp_path / "blocked-dir").exists()
    assert denied_computer.root == tmp_path.resolve()

    _computer, allowed = registry_for(tmp_path, approve=True)
    created = parsed(allowed, "filesystem_write", {"path": "new.txt", "content": "created"})
    overwritten = parsed(
        allowed,
        "filesystem_write",
        {"path": "legacy.txt", "content": "résumé", "overwrite": True},
    )

    assert created["created"] is True
    assert overwritten["encoding"] == "cp1252"
    assert existing.read_text(encoding="cp1252") == "résumé"


def test_document_read_txt_and_markdown(tmp_path):
    (tmp_path / "note.txt").write_text("plain text", encoding="utf-8")
    (tmp_path / "guide.md").write_text("# Markdown guide", encoding="utf-8")
    (tmp_path / "data.json").write_text('{"answer": 42}', encoding="utf-8")
    (tmp_path / "table.csv").write_text("name,value\nTieru,8\n", encoding="utf-8")
    _computer, registry = registry_for(tmp_path)

    text = parsed(registry, "document_read", {"path": "note.txt"})
    markdown = parsed(registry, "document_read", {"path": "guide.md"})
    json_document = parsed(registry, "document_read", {"path": "data.json"})
    csv_document = parsed(registry, "document_read", {"path": "table.csv"})

    assert text == {
        "path": "note.txt", "type": "txt", "content": "plain text", "truncated": False
    }
    assert markdown["type"] == "md" and "# Markdown guide" in markdown["content"]
    assert '"answer": 42' in json_document["content"]
    assert "Tieru\t8" in csv_document["content"]


@pytest.mark.parametrize("filename", ["fixture.pdf", "fixture.docx", "fixture.pptx", "fixture.xlsx"])
def test_document_read_pdf_and_office_fixtures(filename, tmp_path):
    expected = make_documents(tmp_path)
    _computer, registry = registry_for(tmp_path)

    result = parsed(registry, "document_read", {"path": filename})

    assert result["path"] == filename
    assert result["type"] == Path(filename).suffix[1:]
    assert expected[filename] in result["content"]
    assert result["truncated"] is False


def test_unsupported_and_image_only_documents_fail_safely(tmp_path):
    from pypdf import PdfWriter

    (tmp_path / "image.png").write_bytes(b"\x89PNG\r\n")
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with (tmp_path / "blank.pdf").open("wb") as handle:
        writer.write(handle)
    _computer, registry = registry_for(tmp_path)

    unsupported = parsed(registry, "document_read", {"path": "image.png"})
    image_only = parsed(registry, "document_read", {"path": "blank.pdf"})

    assert unsupported["error"]["code"] == "tool_execution_error"
    assert "unsupported document type" in unsupported["error"]["message"]
    assert image_only["error"]["code"] == "tool_execution_error"
    assert "image-only/scanned" in image_only["error"]["message"]


def test_shell_command_success_is_bounded_and_redacted(monkeypatch, tmp_path):
    _computer, registry = registry_for(tmp_path)

    success = parsed(registry, "shell_run", {"command": "git --version"})
    assert success["exit_code"] == 0
    assert "git version" in success["stdout"].lower()
    assert success["cwd"] == "."

    def secret_output(*_args, **_kwargs):
        return SimpleNamespace(
            returncode=0,
            stdout=b"api_key=sk-super-secret-value",
            stderr=b"",
        )

    monkeypatch.setattr(subprocess, "run", secret_output)
    redacted = registry.execute("shell_run", {"command": "safe-command"})
    assert "sk-super-secret-value" not in redacted
    assert "[REDACTED SECRET]" in redacted


def test_shell_timeout_is_controlled_and_releases_process(monkeypatch, tmp_path):
    _computer, registry = registry_for(tmp_path)
    calls = []

    def timeout(*args, **kwargs):
        calls.append((args, kwargs))
        raise subprocess.TimeoutExpired(args[0], kwargs["timeout"])

    monkeypatch.setattr(subprocess, "run", timeout)
    result = parsed(
        registry, "shell_run", {"command": "safe-command", "timeout_seconds": 1}
    )

    assert result["error"]["code"] == "tool_execution_error"
    assert "timed out after 1 seconds" in result["error"]["message"]
    assert calls[0][1]["stdin"] is subprocess.DEVNULL
    assert calls[0][1]["shell"] is False


def test_shell_permission_denial_and_destructive_block(tmp_path):
    computer, denied = registry_for(tmp_path, approve=False)
    denied_result = parsed(denied, "shell_run", {"command": "git status"})

    assert denied_result["error"]["code"] == "tool_permission_denied"

    _computer, approved = registry_for(tmp_path, approve=True)
    blocked = parsed(approved, "shell_run", {"command": "git reset --hard"})

    assert blocked["error"]["code"] == "tool_execution_error"
    assert "Git mutation is unavailable" in blocked["error"]["message"]
    assert computer.root == tmp_path.resolve()


def test_document_resources_are_closed_after_read(tmp_path):
    make_documents(tmp_path)
    computer = LocalComputer(tmp_path)

    computer.document_read("fixture.pdf")
    computer.document_read("fixture.xlsx")
    (tmp_path / "fixture.pdf").replace(tmp_path / "moved.pdf")
    (tmp_path / "fixture.xlsx").replace(tmp_path / "moved.xlsx")

    assert (tmp_path / "moved.pdf").exists()
    assert (tmp_path / "moved.xlsx").exists()


def test_replay_records_filesystem_document_and_shell_actions(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "answer.md").write_text("M8 replay fixture answer.", encoding="utf-8")
    gate = response([text_block('{"retrieve": false, "query": "", "reason": "test"}')])
    client = ScriptedClient(
        [
            gate,
            response(
                [tool_block("filesystem_search", {"query": "replay fixture"}, "search")],
                "tool_use",
            ),
            response(
                [tool_block("filesystem_read", {"path": "answer.md"}, "read")], "tool_use"
            ),
            response(
                [tool_block("document_read", {"path": "answer.md"}, "document")], "tool_use"
            ),
            response(
                [tool_block("shell_run", {"command": "git --version"}, "shell")], "tool_use"
            ),
            response([text_block("The fixture answer was read and git is available.")]),
            response([text_block("The local fixture says M8 replay fixture answer; git is available.")]),
        ]
    )
    app = make_waku(tmp_path / "home", client=client)

    result = app.respond("Find and read the replay fixture, then check git.")
    events = app.replay.get_events(result.run_id)
    completed = {
        event["tool"] for event in events if event["event_type"] == "tool_completed"
    }

    assert {
        "filesystem_search", "filesystem_read", "document_read", "shell_run"
    } <= completed
    assert any(event["event_type"] == "trust_decision" for event in events)
    app.close()
